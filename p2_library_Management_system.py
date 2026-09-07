import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment


# =========================================================
# BOOK CLASS
# =========================================================

class Book:

    def __init__(
        self,
        book_id,
        title,
        author,
        total_copies,
        available_copies=None
    ):
        self.book_id = book_id
        self.title = title
        self.author = author
        self.total_copies = total_copies

        if available_copies is None:
            self.available_copies = total_copies
        else:
            self.available_copies = available_copies

    @property
    def issued_copies(self):
        return self.total_copies - self.available_copies


# =========================================================
# LIBRARY CLASS
# =========================================================

class Library:

    def __init__(self):
        self.books = []
        self.file_path = os.path.join(
            os.path.dirname(os.path.abspath(__file__)),
            "Library_Data.xlsx"
        )
        self.load_books()

    # =====================================================
    # CREATE EXCEL FILE
    # =====================================================
    def create_file(self):
        if os.path.exists(self.file_path):
            return

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Library"

        sheet.append([
            "Book ID",
            "Book Title",
            "Author",
            "Total Copies",
            "Available Copies",
            "Issued Copies"
        ])

        workbook.save(self.file_path)

    # =====================================================
    # FORMAT EXCEL
    # =====================================================
    def format_excel(self, sheet):
        header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        header_font = Font(bold=True, color="FFFFFF")
        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        # Format Header Row (Row 1)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Format Data Rows (Row 2 onwards)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")

        # Column widths
        widths = {"A": 15, "B": 35, "C": 25, "D": 15, "E": 18, "F": 15}
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

        # Freeze header
        sheet.freeze_panes = "A2"

        # Filter
        if sheet.max_row > 1:
            sheet.auto_filter.ref = sheet.dimensions

    # =====================================================
    # SAVE BOOKS
    # =====================================================
    def save_books(self):
        self.create_file()
        workbook = load_workbook(self.file_path)
        sheet = workbook["Library"]

        # Remove old records safely
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row)

        # Write current records
        for book in self.books:
            sheet.append([
                book.book_id,
                book.title,
                book.author,
                book.total_copies,
                book.available_copies,
                book.issued_copies
            ])

        self.format_excel(sheet)
        workbook.save(self.file_path)

    # =====================================================
    # LOAD BOOKS
    # =====================================================
    def load_books(self):
        self.create_file()
        workbook = load_workbook(self.file_path, data_only=True)
        sheet = workbook["Library"]

        # Clear memory array before reloading to avoid double-appending
        self.books = []

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue

            book = Book(
                str(row[0]),
                str(row[1]),
                str(row[2]),
                int(row[3]),
                int(row[4])
            )
            self.books.append(book)

    # =====================================================
    # FIND BOOK
    # =====================================================
    def find_book(self, book_id):
        for book in self.books:
            if book.book_id.lower() == book_id.lower():
                return book
        return None

    # =====================================================
    # 1. ADD BOOK
    # =====================================================
    def add_book(self):
        print("\n========== ADD BOOK ==========")
        book_id = input("Enter Book ID: ").strip()

        if not book_id:
            print("Book ID cannot be empty.")
            return

        if self.find_book(book_id):
            print("\nThis Book ID already exists.")
            print("Use 'Manage Copies' to increase the stock.")
            return

        title = input("Enter Book Title: ").strip()
        author = input("Enter Author Name: ").strip()

        if not title or not author:
            print("\nAll fields are required.")
            return

        while True:
            try:
                copies = int(input("How many copies should be stored? "))
                if copies > 0:
                    break
                print("Copies must be greater than 0.")
            except ValueError:
                print("Please enter a valid number.")

        book = Book(book_id, title, author, copies)
        self.books.append(book)
        self.save_books()
        print("\nBook added successfully!")

    # =====================================================
    # 2. VIEW ALL BOOKS
    # =====================================================
    def view_all_books(self):
        print("\n========== VIEW ALL BOOKS ==========")
        if not self.books:
            print("No books available in the library.")
            return

        for index, book in enumerate(self.books, start=1):
            print(f"\n--- Book {index} ---")
            print(f"Book ID: {book.book_id}")
            print(f"Book Title: {book.title}")
            print(f"Author: {book.author}")
            print(f"Total Copies: {book.total_copies}")
            print(f"Available Copies: {book.available_copies}")
            print(f"Issued Copies: {book.issued_copies}")
        print("\n------------------------------------")

    # =====================================================
    # 3. SEARCH BOOK
    # =====================================================
    def search_book(self):
        print("\n========== SEARCH BOOK ==========")
        print("1. Search by Book ID")
        print("2. Search by Book Title")
        choice = input("\nEnter your choice: ").strip()

        if choice == "1":
            book_id = input("Enter Book ID: ").strip()
            if not book_id:
                print("Book ID cannot be empty.")
                return

            book = self.find_book(book_id)
            if book:
                print("\nBook Found:")
                print(f"Book ID: {book.book_id}")
                print(f"Book Title: {book.title}")
                print(f"Author: {book.author}")
                print(f"Total Copies: {book.total_copies}")
                print(f"Available Copies: {book.available_copies}")
                print(f"Issued Copies: {book.issued_copies}")
            else:
                print("\nNo book found with the given Book ID.")

        elif choice == "2":
            title = input("Enter Book Title: ").strip()
            if not title:
                print("Book Title cannot be empty.")
                return

            matching_books = [
                b for b in self.books
                if title.lower() in b.title.lower()
            ]

            if matching_books:
                print(f"\nFound {len(matching_books)} matching book(s):")
                for index, book in enumerate(matching_books, start=1):
                    print(f"\n--- Result {index} ---")
                    print(f"Book ID: {book.book_id}")
                    print(f"Book Title: {book.title}")
                    print(f"Author: {book.author}")
                    print(f"Total Copies: {book.total_copies}")
                    print(f"Available Copies: {book.available_copies}")
                    print(f"Issued Copies: {book.issued_copies}")
            else:
                print("\nNo books found matching the given title.")

        else:
            print("Invalid choice selection.")

    # =====================================================
    # 4. ISSUE BOOK
    # =====================================================
    def issue_book(self):
        print("\n========== ISSUE BOOK ==========")
        book_id = input("Enter Book ID: ").strip()

        if not book_id:
            print("Book ID cannot be empty.")
            return

        book = self.find_book(book_id)
        if not book:
            print("\nBook not found. Please enter a valid Book ID.")
            return

        if book.available_copies <= 0:
            print(f"\nNo copies available for '{book.title}'. All copies are currently issued.")
            return

        book.available_copies -= 1
        self.save_books()
        print(f"\nBook '{book.title}' issued successfully!")
        print(f"Available Copies: {book.available_copies} | Issued Copies: {book.issued_copies}")

    # =====================================================
    # 5. RETURN BOOK
    # =====================================================
    def return_book(self):
        print("\n========== RETURN BOOK ==========")
        book_id = input("Enter Book ID: ").strip()

        if not book_id:
            print("Book ID cannot be empty.")
            return

        book = self.find_book(book_id)
        if not book:
            print("\nBook not found. Please enter a valid Book ID.")
            return

        if book.issued_copies <= 0:
            print(f"\nCannot return. No copies of '{book.title}' are currently issued.")
            return

        book.available_copies += 1
        self.save_books()
        print(f"\nBook '{book.title}' returned successfully!")
        print(f"Available Copies: {book.available_copies} | Issued Copies: {book.issued_copies}")

    # =====================================================
    # 6. MANAGE COPIES
    # =====================================================
    def manage_copies(self):
        print("\n========== MANAGE COPIES ==========")
        book_id = input("Enter Book ID: ").strip()

        if not book_id:
            print("Book ID cannot be empty.")
            return

        book = self.find_book(book_id)
        if not book:
            print("\nBook not found.")
            return

        print(f"\nBook: {book.title}")
        print(f"Current Total Copies: {book.total_copies}")
        print(f"Available Copies: {book.available_copies}")
        print(f"Issued Copies: {book.issued_copies}")

        print("\n1. Increase Copies")
        print("2. Decrease Copies")
        choice = input("\nEnter your choice: ").strip()

        if choice == "1":
            try:
                amount = int(input("How many copies to add? "))
                if amount <= 0:
                    print("Amount must be greater than 0.")
                    return
                book.total_copies += amount
                book.available_copies += amount
                self.save_books()
                print(f"\nSuccess! Total copies updated to: {book.total_copies}")
            except ValueError:
                print("Please enter a valid number.")

        elif choice == "2":
            try:
                amount = int(input("How many copies to remove? "))
                if amount <= 0:
                    print("Amount must be greater than 0.")
                    return
                if amount > book.available_copies:
                    print(f"\nCannot decrease. Only {book.available_copies} copies are on the shelf.")
                    return
                book.total_copies -= amount
                book.available_copies -= amount
                self.save_books()
                print(f"\nSuccess! Total copies reduced to: {book.total_copies}")
            except ValueError:
                print("Please enter a valid number.")
        else:
            print("Invalid choice selection.")


# =========================================================
# APPLICATION CONSOLE RUNNER
# =========================================================
if __name__ == "__main__":
    library = Library()
    while True:
        print("\n==================================")
        print("    LIBRARY MANAGEMENT SYSTEM     ")
        print("==================================")
        print("1. Add New Book")
        print("2. View All Books")
        print("3. Search Book")
        print("4. Issue Book")
        print("5. Return Book")
        print("6. Manage Copies")
        print("7. Exit System")

        menu_choice = input("\nSelect an option: ").strip()
        if menu_choice == "1":
            library.add_book()
        elif menu_choice == "2":
            library.view_all_books()
        elif menu_choice == "3":
            library.search_book()
        elif menu_choice == "4":
            library.issue_book()
        elif menu_choice == "5":
            library.return_book()
        elif menu_choice == "6":
            library.manage_copies()
        elif menu_choice == "7":
            print("\nExiting system. Excel database updated.")
            break
        else:
            print("Invalid selection. Try again.")